"""Bounded offline DOCX/PDF renderer. No HTML, shell, remote images or templates."""
import base64
import io
import json
import os
from pathlib import Path
import sys
from xml.sax.saxutils import escape

FONT_DIR = Path(__file__).resolve().parent / "document_fonts"


class UnsupportedCharacters(ValueError):
    def __init__(self, characters):
        super().__init__("Configured font does not cover this text.")
        self.characters = characters


def render_document(spec):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from pypdf import PdfReader

    # Bundled DejaVu Sans (Latin incl. Central/Eastern European, Greek,
    # Cyrillic, arrows, check marks). Deployments can opt into another
    # licensed font pair, e.g. a CJK font, via environment variables.
    regular = os.getenv("AGENT_DOCUMENT_FONT", str(FONT_DIR / "DejaVuSans.ttf"))
    bold = os.getenv("AGENT_DOCUMENT_FONT_BOLD", str(Path(regular).with_name("DejaVuSans-Bold.ttf")))
    if not Path(regular).is_file() or not Path(bold).is_file():
        regular, bold = str(FONT_DIR / "DejaVuSans.ttf"), str(FONT_DIR / "DejaVuSans-Bold.ttf")
    pdfmetrics.registerFont(TTFont("Consens", regular))
    pdfmetrics.registerFont(TTFont("ConsensBold", bold))
    pdfmetrics.registerFontFamily("Consens", normal="Consens", bold="ConsensBold", italic="Consens", boldItalic="ConsensBold")
    # Only the PDF embeds this font; DOCX text is rendered by the reader's
    # fonts. Report the exact characters so the model can rewrite them.
    content = json.dumps(spec, ensure_ascii=False)
    missing = set()
    for font in ("Consens", "ConsensBold"):
        glyphs = pdfmetrics.getFont(font).face.charToGlyph
        missing.update(c for c in content if ord(c) >= 32 and ord(c) not in glyphs)
    if missing:
        raise UnsupportedCharacters("".join(sorted(missing))[:40])

    doc = Document()
    page = doc.sections[0]
    page.page_width, page.page_height = Inches(8.27), Inches(11.69)
    page.top_margin = page.bottom_margin = Inches(.7)
    page.left_margin = page.right_margin = Inches(.75)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(22, 30, 42)
    doc.styles["Normal"].font.size = Pt(10)
    doc.styles["Normal"].paragraph_format.space_after = Pt(7)
    doc.styles["Normal"].paragraph_format.widow_control = True
    doc.core_properties.title = spec["title"]
    doc.core_properties.author = "Consens"
    styles = getSampleStyleSheet()
    for name in ("Normal", "Title", "Heading1", "Heading2"):
        styles[name].fontName = "ConsensBold" if name != "Normal" else "Consens"
        styles[name].fontSize = {"Normal": 10, "Title": 23, "Heading1": 15, "Heading2": 12}[name]
        styles[name].leading = styles[name].fontSize * 1.35
        styles[name].spaceAfter = 8
        styles[name].splitLongWords = True
    cellstyle = ParagraphStyle("Cell", parent=styles["Normal"], fontSize=8, leading=11, spaceAfter=0)
    story = []
    def paragraph(text, style="Normal"):
        doc.add_paragraph(text, style={"Heading1": "Heading 1", "Title": "Title"}.get(style, style))
        story.append(Paragraph(escape(text).replace("\n", "<br/>"), styles[style]))
    paragraph(spec["title"], "Title")
    if spec["summary"]:
        paragraph(spec["summary"])
    for index, section in enumerate(spec["sections"], 1):
        paragraph(f"{index}. {section['heading']}", "Heading1")
        for text in section["paragraphs"]:
            paragraph(text)
        if section["table"]:
            table = section["table"]
            rows = [table["headers"], *table["rows"]]
            dt = doc.add_table(rows=1, cols=len(table["headers"]))
            dt.style = "Table Grid"
            header = OxmlElement("w:tblHeader"); dt.rows[0]._tr.get_or_add_trPr().append(header)
            for i, row in enumerate(rows):
                cells = dt.rows[0].cells if i == 0 else dt.add_row().cells
                for j, text in enumerate(row):
                    cells[j].text = text
                    if i == 0:
                        for run in cells[j].paragraphs[0].runs:
                            run.bold = True
                        shading = OxmlElement("w:shd"); shading.set(qn("w:fill"), "E8EDF3"); cells[j]._tc.get_or_add_tcPr().append(shading)
            pt = Table([[Paragraph(escape(cell).replace("\n", "<br/>"), cellstyle) for cell in row] for row in rows],
                colWidths=[487 / len(table["headers"])] * len(table["headers"]), repeatRows=1, hAlign="LEFT")
            pt.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EDF3")),
                ("GRID", (0, 0), (-1, -1), .4, colors.HexColor("#B9C3CF")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
            story.extend([pt, Spacer(1, 12)])
    for heading, key in [("Uncertainties and limits", "uncertainties"), ("Differing assessments", "differing_views")]:
        if spec[key]:
            paragraph(heading, "Heading1")
            for text in spec[key]:
                paragraph(text)
    if spec["sources"]:
        paragraph("Sources", "Heading1")
        for i, source in enumerate(spec["sources"], 1):
            paragraph(f"[{i}] {source['label']}" + (f" — {source['locator']}" if source["locator"] else "") +
                (f" — {source['url']} (model-supplied reference; verify against the original)" if source["url"] else f" — file {source['file_id']}"))
    out_doc = io.BytesIO(); doc.save(out_doc)
    out_pdf = io.BytesIO()
    def footer(canvas, pdf):
        canvas.setFont("Consens", 8); canvas.drawRightString(540, 28, str(pdf.page))
    SimpleDocTemplate(out_pdf, pagesize=(595.28, 841.89), leftMargin=54, rightMargin=54,
        topMargin=50, bottomMargin=50, title=spec["title"], author="Consens").build(story, onFirstPage=footer, onLaterPages=footer)
    # Reopen both actual binary formats before publishing anything.
    reopened = Document(io.BytesIO(out_doc.getvalue()))
    pdf = PdfReader(io.BytesIO(out_pdf.getvalue()), strict=True)
    if not reopened.paragraphs or not 1 <= len(pdf.pages) <= 100:
        raise ValueError("Invalid document output.")
    if spec["title"] not in reopened.paragraphs[0].text or not pdf.pages[0].extract_text():
        raise ValueError("Document content is missing.")
    return {"docx": out_doc.getvalue(), "pdf": out_pdf.getvalue()}


if __name__ == "__main__":
    import resource
    resource.setrlimit(resource.RLIMIT_CPU, (20, 20))
    resource.setrlimit(resource.RLIMIT_AS, (1024 * 1024 * 1024, 1024 * 1024 * 1024))
    # Lightweight import only: no Firestore, Firebase Admin or credentials here.
    from app.services.agent_document_spec import DocumentSpec
    raw = sys.stdin.buffer.read(50_001)
    if len(raw) > 50_000:
        raise ValueError("Document input too large.")
    spec = DocumentSpec.model_validate_json(raw).model_dump()
    try:
        rendered = render_document(spec)
    except UnsupportedCharacters as exc:
        print(json.dumps({"error": "unsupported_characters", "characters": exc.characters}))
        sys.exit(0)
    print(json.dumps({key: base64.b64encode(value).decode() for key, value in rendered.items()}))
