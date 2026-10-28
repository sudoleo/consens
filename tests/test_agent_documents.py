import io
from types import SimpleNamespace
import pytest
from docx import Document
from pypdf import PdfReader
from app.services.agent_documents import CreateDocument, DocumentTools, ReadDocument, ReviseDocument
from app.services.agent_files import FileContext, FileUnavailable
from app.services.llm.provider_runtime import ProviderCancellation
from test_agent_files import setup, upload


def sample():
    return {"document": {"title": "Decision brief", "summary": "Evaluate the two vendor offers.",
        "sections": [{"heading": "Comparison", "paragraphs": ["Vendor A is cheaper; vendor B ships sooner."],
            "table": {"headers": ["Vendor", "Price", "Delivery"], "rows": [["A", "42 EUR", "2 weeks"], ["B", "55 EUR", "1 week"]]}},
            {"heading": "Action plan", "paragraphs": ["Confirm delivery with procurement."]}],
        "uncertainties": ["Tax treatment is unconfirmed."], "differing_views": ["Model B prioritizes delivery over price."]}}


def service(files, chat, turn="first", uid="owner"):
    return DocumentTools(SimpleNamespace(file_context=FileContext(files, uid, chat, []), uid=uid, chat_id=chat, turn_id=turn))


def test_real_docx_pdf_version_and_saved_provenance(setup):
    files, chat = setup
    source = upload(files, chat)
    args = sample(); args["document"]["sources"] = [{"label": "Vendor quote", "file_id": source["id"], "locator": "lines 1-2"}]
    docs = service(files, chat)
    created = docs.create(CreateDocument.model_validate(args), cancellation=ProviderCancellation())
    assert len(created["files"]) == 2
    for f in created["files"]:
        meta, raw = files.download("owner", chat, f["id"])
        assert meta["turn_id"] == "first" and meta["source_file_ids"] == [source["id"]]
        if f["mime"] == "application/pdf":
            pdf = PdfReader(io.BytesIO(raw), strict=True)
            text = "\n".join(p.extract_text() for p in pdf.pages)
            assert "Tax treatment" in text and "Model B prioritizes" in text and "42 EUR" in text
        else:
            document = Document(io.BytesIO(raw))
            assert document.tables[0].cell(1, 1).text == "42 EUR"
    saved = docs.read(ReadDocument(document_id=created["document_id"]), cancellation=ProviderCancellation())
    assert saved["sources"][0]["sha256"] == source["sha256"]
    assert saved["content_hash"] == created["content_hash"]
    repeat = docs.create(CreateDocument.model_validate(args), cancellation=ProviderCancellation())
    assert repeat["files"] == created["files"] and len(files.list("owner", chat)) == 3
    revision = service(files, chat, "followup").revise(ReviseDocument(document_id=created["document_id"], version=1,
        section_number=2, replacement={"heading": "Revised plan", "paragraphs": ["Contact vendor B on Monday."]}, change_summary="Revised section two"), cancellation=ProviderCancellation())
    assert revision["version"] == 2 and revision["parent_version"] == 1
    latest = docs.read(ReadDocument(document_id=created["document_id"]), cancellation=ProviderCancellation())
    original = docs.read(ReadDocument(document_id=created["document_id"], version=1), cancellation=ProviderCancellation())
    assert original["content"]["sections"][1]["heading"] == "Action plan"
    assert latest["content"]["sections"][1]["heading"] == "Revised plan"
    assert latest["content"]["sections"][0] == original["content"]["sections"][0]
    assert latest["content"]["uncertainties"] == original["content"]["uncertainties"]
    with pytest.raises(FileUnavailable, match="newer"):
        docs.revise(ReviseDocument(document_id=created["document_id"], version=1, section_number=2,
            replacement={"heading": "Stale", "paragraphs": []}, change_summary="stale edit"), cancellation=ProviderCancellation())


def test_ownership_invalid_sources_and_storage_failure(setup, monkeypatch):
    from app.services.chat_store import ChatNotFound
    files, chat = setup
    docs = service(files, chat)
    created = docs.create(CreateDocument.model_validate(sample()), cancellation=ProviderCancellation())
    with pytest.raises(ChatNotFound):
        service(files, chat, uid="other").read(ReadDocument(document_id=created["document_id"]), cancellation=ProviderCancellation())
    other_chat = files.chats.create_chat("owner", execution_mode="agent")["id"]
    with pytest.raises(FileUnavailable):
        service(files, other_chat).read(ReadDocument(document_id=created["document_id"]), cancellation=ProviderCancellation())
    args = sample(); args["document"]["sources"] = [{"label": "Missing", "file_id": "a" * 32}]
    with pytest.raises(FileUnavailable):
        docs.create(CreateDocument.model_validate(args), cancellation=ProviderCancellation())
    real_put = files.objects.put; calls = []
    def fail_second(key, raw, mime):
        calls.append(key)
        if len(calls) == 2:
            raise OSError("Storage unavailable")
        real_put(key, raw, mime)
    monkeypatch.setattr(files.objects, "put", fail_second)
    before = files.list("owner", chat)
    with pytest.raises(OSError):
        service(files, chat, "failed").create(CreateDocument.model_validate(sample()), cancellation=ProviderCancellation())
    assert files.list("owner", chat) == before
    with pytest.raises(FileUnavailable):
        files.objects.get(calls[0])


def test_limits_unsupported_glyphs_and_cancellation(setup):
    from pydantic import ValidationError
    from app.services.llm.provider_runtime import ProviderCancelled
    files, chat = setup
    args = sample(); args["document"]["sections"][0]["paragraphs"] = ["a" * 40_000]
    with pytest.raises(ValidationError): CreateDocument.model_validate(args)
    args = sample(); args["document"]["sources"] = [{"label": "Bad", "url": "javascript:alert(1)"}]
    with pytest.raises(ValidationError): CreateDocument.model_validate(args)
    args = sample(); args["document"]["title"] = "Unsupported \U0001f999"
    with pytest.raises(FileUnavailable, match="cannot display these characters"):
        service(files, chat).create(CreateDocument.model_validate(args), cancellation=ProviderCancellation())
    cancelled = ProviderCancellation(); cancelled.cancel()
    with pytest.raises(ProviderCancelled):
        service(files, chat).create(CreateDocument.model_validate(sample()), cancellation=cancelled)
    assert not files.list("owner", chat)


def test_the_answer_is_written_from_the_actual_document_results(setup):
    """The agent writes the answer itself after create_document: the step
    that writes it reads the real result (title, saved files) and the rule
    never to invent links to the cards."""
    import json
    from test_agent_comparison import Script
    from test_agent_runs import UID, pending
    from app.services.agent_comparison import comparison_selection, review_is_bound
    from app.services.agent_delegation import DelegationLoop
    from app.services.agent_delegation_config import defaults
    from app.services.agent_policy import AgentPolicy
    from app.services.agent_runs import AgentRunStore
    from app.services.llm.agent_client import measured_usage, resolve_agent_model
    files, _ = setup
    store = AgentRunStore(files.db)
    chat, turn = pending(store)
    script = Script()
    base = type(script.factory())
    steps = {}

    class Workflow(base):
        def stream(self, **kwargs):
            if self.step_id.startswith("completion:"):
                steps[int(self.step_id.split(":")[-1])] = kwargs
            if self.step_id == "completion:1" and kwargs["tools"]:
                self.usage = measured_usage({"prompt_tokens": 50, "completion_tokens": 20, "cost": .0001}, kwargs["model"])
                self.tool_calls = [{"id": "doc", "type": "function", "function": {
                    "name": "create_document", "arguments": json.dumps(sample())}}]
                self.finish_reason = "tool_calls"
                return
            yield from super().stream(**kwargs)

    config = {**defaults(), "enabled": False, "max_searches": 0, "context_chars": 120_000}
    loop = DelegationLoop(store=store, uid=UID, chat_id=chat, turn_id=turn["id"],
        model=resolve_agent_model("claude-haiku-4-5"),
        messages=[{"role": "system", "content": "Answer."}, {"role": "user", "content": "Compare the offers and write a brief."}],
        api_key="test", cancellation=ProviderCancellation(), policy=AgentPolicy.from_config({**config, "enabled": True}),
        delegation_config=config, completion_factory=Workflow, file_context=FileContext(files, UID, chat, []),
        comparison_models=comparison_selection({"anthropic": "claude-haiku-4-5", "openai": "gpt-5.4-mini"}))
    script.loop = loop
    events = list(loop.run())
    saved = store.get_turn(UID, chat, turn["id"])
    assert saved["status"] == "completed" and review_is_bound(saved["agent_review"], saved["consensus"])
    document = loop.documents.results[-1]
    assert document["title"] == "Decision brief" and len(document["files"]) == 2
    # Compare, create the document, write the answer: no further step.
    assert sorted(steps) == [0, 1, 2] and saved["consensus"] == "The first option costs 100."
    writer = steps[2]
    assert "never invent links to them" in writer["messages"][0]["content"]
    result = json.loads(next(m["content"] for m in writer["messages"] if m.get("tool_call_id") == "doc"))
    assert result["document_id"] == document["document_id"] and result["title"] == "Decision brief"
    assert {f["id"] for f in result["files"]} == {f["id"] for f in document["files"]}
    assert any(e.get("type") == "resources" for e in events)


def test_the_file_list_travels_with_the_message_not_the_system_prompt(setup):
    """The chat's file list changes with every upload. In the system prompt it
    rewrote the prompt cache of everything after it, the history included;
    in the message's app context the system prompt stays the same."""
    import base64
    from test_agent_runs import UID, pending
    from app.services.agent_comparison import comparison_selection
    from app.services.agent_delegation import DelegationLoop
    from app.services.agent_delegation_config import defaults
    from app.services.agent_files import UNTRUSTED
    from app.services.agent_policy import AgentPolicy
    from app.services.agent_runs import APP_CONTEXT_CLOSE, AgentRunStore, strip_app_context, with_app_context
    from app.services.llm.agent_client import AgentCompletion, resolve_agent_model
    files, _ = setup
    store = AgentRunStore(files.db)
    chat, turn = pending(store)
    config = {**defaults(), "enabled": False, "max_searches": 0, "context_chars": 120_000}

    def loop():
        return DelegationLoop(store=store, uid=UID, chat_id=chat, turn_id=turn["id"],
            model=resolve_agent_model("claude-haiku-4-5"),
            messages=[{"role": "system", "content": "Answer."},
                      {"role": "user", "content": with_app_context("Compare the offers.", "Current date: today.")}],
            api_key="test", cancellation=ProviderCancellation(), policy=AgentPolicy.from_config({**config, "enabled": True}),
            delegation_config=config, completion_factory=AgentCompletion, file_context=FileContext(files, UID, chat, []),
            comparison_models=comparison_selection({"anthropic": "claude-haiku-4-5", "openai": "gpt-5.4-mini"}))

    files.upload(UID, chat, {"name": "offer.txt", "data": base64.b64encode(b"Price: 42 EUR").decode()})
    first = loop()
    files.upload(UID, chat, {"name": "second.txt", "data": base64.b64encode(b"Price: 55 EUR").decode()})
    second = loop()
    assert first.messages[0] == second.messages[0]
    assert UNTRUSTED in first.messages[0]["content"] and "Files available" not in first.messages[0]["content"]
    question = second.messages[-1]["content"]
    context = question.split(APP_CONTEXT_CLOSE)[0]
    assert "Current date: today." in context and "Files available in this chat:" in context
    assert "offer.txt" in context and "second.txt" in context
    # The user's own words stay what they wrote (memory evidence, read_source links).
    assert strip_app_context(question) == "Compare the offers." == second.answer_conversation[-1]["content"]


def test_chat_deletion_removes_document_versions_and_manifest(setup):
    files, chat = setup
    docs = service(files, chat)
    created = docs.create(CreateDocument.model_validate(sample()), cancellation=ProviderCancellation())
    manifest = docs.ref(created["document_id"])
    assert manifest.collection("versions").document("1").get().exists
    files.chats.delete_chat("owner", chat)
    assert not manifest.get().exists
    assert not manifest.collection("versions").document("1").get().exists
    assert not any("documents" in path for path in files.db.documents)


def test_render_subprocess_does_not_import_firebase():
    import subprocess, sys
    code = ("import sys, app.services.agent_document_render as r, app.services.agent_document_spec; "
            "print('firebase_admin' in sys.modules)")
    out = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True, check=True)
    assert out.stdout.strip() == "False"


def test_common_symbols_and_central_european_text_render(setup):
    files, chat = setup
    args = sample()
    args["document"]["sections"][0]["paragraphs"] = ["Łódź → Győr ✓ ąę Ωmega Жизнь – “quoted” …"]
    created = service(files, chat).create(CreateDocument.model_validate(args), cancellation=ProviderCancellation())
    pdf = next(f for f in created["files"] if f["mime"] == "application/pdf")
    text = PdfReader(io.BytesIO(files.download("owner", chat, pdf["id"])[1])).pages[0].extract_text()
    assert "Győr" in text and "Łódź" in text


def test_unsupported_characters_are_named_in_the_error(setup):
    files, chat = setup
    args = sample()
    args["document"]["sections"][0]["paragraphs"] = ["Chinese: 中文"]
    with pytest.raises(FileUnavailable, match="中"):
        service(files, chat).create(CreateDocument.model_validate(args), cancellation=ProviderCancellation())
    assert not files.list("owner", chat)


def test_control_characters_are_rejected_before_rendering():
    from pydantic import ValidationError
    args = sample()
    args["document"]["summary"] = "bad\x0bvalue"
    with pytest.raises(ValidationError, match="control characters"):
        CreateDocument.model_validate(args)


def test_document_tools_get_a_larger_argument_limit_than_other_tools():
    from pydantic import BaseModel, ConfigDict
    from app.services.agent_tools import ReadOnlyTool, ToolRegistry
    class Args(BaseModel):
        model_config = ConfigDict(extra="forbid", strict=True)
        text: str
    registry = ToolRegistry([ReadOnlyTool("small", "", Args, None), ReadOnlyTool("large", "", Args, None, argument_limit=50_000)],
                            argument_limit=24_000)
    payload = '{"text": "' + "x" * 30_000 + '"}'
    assert registry.argument_limit == 50_000
    assert registry.validate({"function": {"name": "large", "arguments": payload}})[0].name == "large"
    with pytest.raises(ValueError, match="not authorized"):
        registry.validate({"function": {"name": "small", "arguments": payload}})


def test_failed_revision_keeps_the_manifest_title(setup, monkeypatch):
    files, chat = setup
    docs = service(files, chat)
    created = docs.create(CreateDocument.model_validate(sample()), cancellation=ProviderCancellation())
    monkeypatch.setattr("app.services.agent_documents.render", lambda spec: (_ for _ in ()).throw(FileUnavailable("render failed")))
    with pytest.raises(FileUnavailable):
        docs.revise(ReviseDocument(document_id=created["document_id"], version=1, section_number=1,
            replacement={"heading": "X", "paragraphs": []}, change_summary="x"), cancellation=ProviderCancellation())
    manifest = docs.ref(created["document_id"]).get().to_dict()
    assert manifest["title"] == "Decision brief" and manifest["version"] == 1
    assert "writing_until" not in manifest and "operation" not in manifest


def test_file_list_is_chronological_and_carries_document_titles(setup):
    files, chat = setup
    source = upload(files, chat)
    docs = service(files, chat)
    created = docs.create(CreateDocument.model_validate(sample()), cancellation=ProviderCancellation())
    service(files, chat, "followup").revise(ReviseDocument(document_id=created["document_id"], version=1, section_number=2,
        replacement={"heading": "Revised plan", "paragraphs": ["Call vendor B."]}, change_summary="Revised"), cancellation=ProviderCancellation())
    listed = files.list("owner", chat)
    assert [f["created_at"] for f in listed] == sorted(f["created_at"] for f in listed)
    assert listed[0]["id"] == source["id"] and "title" not in listed[0]
    documents = [f for f in listed if f.get("kind") == "document"]
    assert [f["version"] for f in documents] == [1, 1, 2, 2]
    assert all(f["title"] == "Decision brief" for f in documents)
    # Versions saved before titles lived on file records fall back to the manifest.
    for f in documents:
        record = files.ref("owner", chat, f["id"])
        data = record.get().to_dict()
        data.pop("title")
        record.set(data)
    assert all(f["title"] == "Decision brief" for f in files.list("owner", chat) if f.get("kind") == "document")


def test_audit_mode_names_no_model_comparisons_for_a_document():
    """Opus in audit mode, 2026-10-28: create_document asked for model
    comparisons it had no tool for. Only comparison mode names them."""
    from types import SimpleNamespace
    from app.services.agent_documents import DocumentTools

    def description(audit):
        tools = object.__new__(DocumentTools)
        tools.loop = SimpleNamespace(comparison=SimpleNamespace(audit_mode=audit))
        return next(tool.description for tool in tools.tools() if tool.name == "create_document")
    assert "comparison" not in description(True).lower() and "Preserve uncertainties" in description(True)
    assert "Complete model comparisons first" in description(False)
